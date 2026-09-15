# -*- coding: utf-8 -*-
"""基于 pandoc 把 11 个公式写成 LaTeX 数学并转成 docx，抽取原生 OMML 对象，注入主文档对应段落。"""
import re, subprocess, os, shutil
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

XG4 = r"D:\TraeNewPaper07-1\Papers2\v18\v18-3\v18-3-3\v18-3-3-2\xg4"

# 11 个公式（LaTeX 显示数学），顺序与论文式(1)~(11)一致
formulas = {
 1: r"S_{it}^{*}=100\times\sum_{j}w_{j}z_{ijt}^{*},\quad \sum_{j}w_{j}=1",
 2: r"Y_{it}=S_{i,t+1}^{*}-S_{it}^{*}",
 3: r"\tau(x)=E[Y(1)-Y(0)\mid X=x]",
 4: r"m(x)=E(Y\mid X=x),\quad e(x)=P(T=1\mid X=x)",
 5: r"\hat{Y}_{i}=Y_{i}-\hat{m}(X_{i}),\quad \tilde{T}_{i}=T_{i}-\hat{e}(X_{i})",
 6: r"\rho_{i}=\hat{Y}_{i}/\tilde{T}_{i},\quad \omega_{i}=\tilde{T}_{i}^{2}",
 7: r"\hat{\tau}=\arg\min_{\tau}\sum_{i=1}^{n}\bigl(\hat{Y}_{i}-\tau(X_{i})\tilde{T}_{i}\bigr)^{2}",
 8: r"\hat{\tau}=\arg\min_{\tau}\sum_{i=1}^{n}\omega_{i}\bigl(\rho_{i}-\tau(X_{i})\bigr)^{2}",
 9: r"\hat{\tau}(x)=B^{-1}\sum_{b=1}^{B}\hat{\tau}_{b}(x)",
 10: r"I_{j}=G_{j}/\sum_{k}G_{k}",
 11: r"\max\sum_{r=1}^{R}I_{r}(\bar{c}_{r}+\hat{\tau}(x_{r}))c_{r},\quad \text{s.t.}\;\sum_{r=1}^{R}c_{r}\le C,\quad 0\le c_{r}\le \bar{c}_{r}",
}

# ---- 1) 生成公式 docx ----
md_lines = []
for k, latex in formulas.items():
    md_lines.append(f"MARKER {k}\n\n$$\n{latex}\n$$\n")
md = "\n".join(md_lines)
md_path = os.path.join(XG4, "code", "_formulas.md")
open(md_path, "w", encoding="utf-8").write(md)
mdx_path = os.path.join(XG4, "code", "_formulas.docx")
pandoc = r"C:\Users\LENOVO\AppData\Roaming\TRAE SOLO CN\ModularData\ai-agent\vm\tools\bin\pandoc.cmd"
r = subprocess.run([pandoc, md_path, "-t", "docx", "-o", mdx_path],
                   capture_output=True, text=True)
print("pandoc公式: exit", r.returncode)

# ---- 2) 抽取每个 oMathPara XML ----
fd = Document(mdx_path)
math_oms = {}   # k -> oMathPara xml string
cur = None
for p in fd.paragraphs:
    t = (p.text or "").strip()
    m = re.match(r"^MARKER (\d+)$", t)
    if m:
        cur = int(m.group(1))
        continue
    # 找该段内第一个 m:oMathPara
    omps = p._p.findall('.//'+qn('m:oMathPara')) or (p._p.findall('.//'+qn('m:oMath')))
    if cur and omps:
        math_oms[cur] = etree.tostring(omps[0], pretty_print=False).decode('utf-8')
        cur = None
print("抽取到公式:", sorted(math_oms.keys()))

# ---- 3) 注入主文档 ----
main_path = os.path.join(XG4, "论文v18-3-3-修订版.docx")
# 先用 pandoc 从定稿 HTML 重建基线，避免上次误注入污染
html_src = os.path.join(XG4, "论文v18-3-3-修订版.html")
subprocess.run([pandoc, html_src, "-f", "html", "-t", "docx", "-o", main_path],
               capture_output=True, text=True)
d = Document(main_path)
NSM = "http://schemas.openxmlformats.org/officeDocument/2006/math"
CJK = re.compile(r'[\u4e00-\u9fff]')

def formula_key(txt):
    """仅当段落为‘独占一行的公式段’时才返回式号：以（k）结尾、且去掉式号后不含中文。"""
    s = txt.strip()
    m = re.match(r'^(.*?)（(\d{1,2})）$', s, re.S)
    if not m:
        return None
    body = m.group(1).strip()
    k = int(m.group(2))
    if 1 <= k <= 11 and CJK.search(body) is None and 3 <= len(body) <= 160:
        return k
    return None

injected = []
for p in d.paragraphs:
    txt = p.text or ""
    k = formula_key(txt)
    if k is None:
        continue
    hit = k
    oms = math_oms.get(hit)
    if not oms:
        print("缺式%d的OMML" % hit); continue
    # 清空该段旧 runs（仅文字、保留段落属性）
    for r in list(p._p.findall(qn('w:r'))):
        p._p.remove(r)
    # 插入 OMML（去掉其外层 oMathPara 的命名空间冲突，添加根命名空间）
    xml = oms
    if xml.startswith("<m:oMath>"):
        xml = f'<m:oMathPara xmlns:m="{NSM}">{xml}</m:oMathPara>'
    elif not xml.startswith("<m:oMathPara"):
        xml = f'<m:oMathPara xmlns:m="{NSM}">{xml}</m:oMathPara>'
    else:
        xml = xml.replace('<m:oMathPara>', f'<m:oMathPara xmlns:m="{NSM}">', 1)
    try:
        node = etree.fromstring(xml.encode('utf-8'))
    except Exception as e:
        print("式%d OMML解析失败: %s" % (hit, e)); continue
    p._p.append(node)
    # 右侧补式号（正常文字，右对齐由段落属性保留）
    rn = etree.SubElement(p._p, qn('w:r'))
    tn = etree.SubElement(rn, qn('w:t'))
    tn.text = f"　（{hit}）"
    injected.append(hit)

d.save(main_path)
print("注入成功:", sorted(injected))
shutil.rmtree(XG4 + r"\code\_tmp.md", ignore_errors=True)