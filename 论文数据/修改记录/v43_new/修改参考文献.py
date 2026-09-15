"""
修改v43.docx：参考文献改为知网优秀论文，重新标注引用，右上角角标
"""
import os, re, copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

SRC = r'd:\TraePP06\论文数据\修改记录\v43\基于LightGBM的零售终端时序预测与预警研究_v43.docx'
OUT = r'd:\TraePP06\论文数据\修改记录\v43_new\基于LightGBM的零售终端时序预测与预警研究_v43_new.docx'

doc = Document(SRC)

# ============ 新参考文献列表（仅知网优秀论文5篇） ============
NEW_REFS = [
    '[1] 姜思明, 刘荣, 谭升达, 等. 基于零售终端数据治理的卷烟市场状态监测研究[J/OL]. 烟草科技, 2026.',
    '[2] 李天举, 谢志峰, 张侃弘, 等. 基于集成学习的烟草异常数据挖掘研究与应用[J]. 计算机技术与发展, 2020, 30(11): 128-135.',
    '[3] 刘忠华, 卢鑫, 梅文强, 等. 基于LightGBM模型的中国成人吸烟行为研究[J]. 现代信息科技, 2024, 8(7): 128-136.',
    '[4] 葛娜, 孙连英, 赵平, 等. 基于ARIMA时间序列模型的销售量预测分析[J]. 北京联合大学学报, 2018, 32(4): 27-33.',
    '[5] 刘海涛. 基于用户行为和智能算法的烟草精准营销系统设计[J]. 国外电子测量技术, 2026, 45(2): 241-246.',
]

# ============ 辅助函数 ============
def get_para_text(p):
    return ''.join(run.text for run in p.runs)

def set_para_text(p, text):
    if not p.runs:
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run._element.getparent().remove(run._element)

def rebuild_para_with_superscript_refs(p, cf='宋体', ef='Times New Roman', fs=Pt(14), bd=False):
    """重建段落：[N]格式的参考文献标记变为上标"""
    full_text = get_para_text(p)
    if not full_text:
        return

    # 清除所有现有run
    for run in list(p.runs):
        run._element.getparent().remove(run._element)

    # 按参考文献标记分割文本
    pattern = r'(\[\d+(?:-\d+)?\])'
    parts = re.split(pattern, full_text)

    for part in parts:
        if not part:
            continue
        if re.match(pattern, part):
            # 参考文献标记 → 上标
            run = p.add_run(part)
            run.font.size = Pt(9)
            run.font.name = ef
            rpr = run._element.get_or_add_rPr()
            vert = OxmlElement('w:vertAlign')
            vert.set(qn('w:val'), 'superscript')
            rpr.append(vert)
            rf_e = OxmlElement('w:rFonts')
            rf_e.set(qn('w:eastAsia'), cf)
            rpr.insert(0, rf_e)
        else:
            # 正常文本
            run = p.add_run(part)
            run.font.size = fs
            run.font.name = ef
            run.bold = bd
            rpr = run._element.get_or_add_rPr()
            rf_e = OxmlElement('w:rFonts')
            rf_e.set(qn('w:eastAsia'), cf)
            rpr.insert(0, rf_e)

# ============ 第一步：修改正文中的引用标注 ============
for p in doc.paragraphs:
    text = get_para_text(p)

    # 姜思明等[2]→[1]
    if '姜思明等[2]' in text:
        text = text.replace('姜思明等[2]', '姜思明等[1]')

    # 李天举等[3]→[2]
    if '李天举等[3]' in text:
        text = text.replace('李天举等[3]', '李天举等[2]')

    # Grinsztajn+Borisov → 刘忠华[3]
    if 'Grinsztajn' in text and 'Borisov' in text:
        # 多种可能的文本格式
        old_patterns = [
            'Grinsztajn等[4]和Borisov等[5]证实GBDT类方法在表格数据建模上优于深度学习方法。',
            'Grinsztajn等[4]和Borisov等[5]证实GBDT类方法在表格数据建模上优于深度学习方法',
        ]
        new_text = '刘忠华等[3]基于LightGBM模型验证了梯度提升决策树在分类预测中的有效性。'
        for old in old_patterns:
            if old in text:
                text = text.replace(old, new_text)
                break
        # 如果上面没匹配到，尝试更宽泛的匹配
        if 'Grinsztajn' in text:
            # 用正则替换整句
            text = re.sub(
                r'Grinsztajn等\[\d+\]和Borisov等\[\d+\][^。]*。',
                '刘忠华等[3]基于LightGBM模型验证了梯度提升决策树在分类预测中的有效性。',
                text
            )

    # 随机森林[6]→[3]
    if '[6]' in text and ('随机森林' in text or 'RF' in text or 'Random Forest' in text):
        text = text.replace('[6]', '[3]')

    # LightGBM[6]→[3] (在模型描述段落中)
    if 'LightGBM（Light Gradient Boosting Machine）[6]' in text:
        text = text.replace('LightGBM（Light Gradient Boosting Machine）[6]',
                           'LightGBM（Light Gradient Boosting Machine）[3]')

    # 已有研究[2-3]→[1-2]
    if '已有研究[2-3]' in text:
        text = text.replace('已有研究[2-3]', '已有研究[1-2]')

    # 如果文本有变化，更新段落
    if text != get_para_text(p):
        set_para_text(p, text)

# ============ 第二步：替换参考文献部分 ============
# 找到参考文献段落
ref_title_para = None
ref_paras = []  # 参考文献条目段落

in_refs = False
for p in doc.paragraphs:
    text = get_para_text(p).strip()
    if text == '参考文献':
        ref_title_para = p
        in_refs = True
        continue
    if in_refs:
        if re.match(r'\[\d+\]', text):
            ref_paras.append(p)
        else:
            break

print(f'找到参考文献标题: {ref_title_para is not None}')
print(f'找到参考文献条目: {len(ref_paras)} 条')

# 替换前N条参考文献内容
for j in range(min(len(ref_paras), len(NEW_REFS))):
    set_para_text(ref_paras[j], NEW_REFS[j])

# 删除多余的参考文献段落
if len(ref_paras) > len(NEW_REFS):
    for j in range(len(NEW_REFS), len(ref_paras)):
        para_elem = ref_paras[j]._element
        para_elem.getparent().remove(para_elem)
    print(f'删除了 {len(ref_paras) - len(NEW_REFS)} 条多余参考文献')

# ============ 第三步：给所有正文中的[N]引用加上标格式 ============
new_ref_prefixes = [ref.split(']')[0] + ']' for ref in NEW_REFS]

for p in doc.paragraphs:
    text = get_para_text(p)
    if not text:
        continue

    # 跳过参考文献条目
    stripped = text.strip()
    if any(stripped.startswith(prefix) for prefix in new_ref_prefixes):
        continue

    # 跳过参考文献标题
    if stripped == '参考文献':
        continue

    # 跳过作者简介
    if stripped.startswith('作者简介'):
        continue

    # 检查是否包含[N]格式的引用
    if re.search(r'\[\d+(?:-\d+)?\]', text):
        rebuild_para_with_superscript_refs(p)

# ============ 保存 ============
doc.save(OUT)
print(f'\n[OK] 论文已修改并保存: {OUT}')

# ============ 验证 ============
doc2 = Document(OUT)
print('\n===== 修改后引用相关段落 =====')
for i, p in enumerate(doc2.paragraphs):
    text = p.text.strip()
    if text and re.search(r'\[\d+', text):
        # 检查是否是参考文献条目
        if re.match(r'\[\d+\]', text) and len(text) < 200:
            continue  # 跳过参考文献条目，单独打印
        print(f'[{i}] {text[:150]}')

print('\n===== 参考文献部分 =====')
in_refs = False
for p in doc2.paragraphs:
    text = p.text.strip()
    if text == '参考文献':
        in_refs = True
        print(text)
    elif in_refs:
        if re.match(r'\[\d+\]', text):
            print(text)
        else:
            break

# 检查是否还有英文参考文献或旧引用
print('\n===== 残留检查 =====')
has_grinsztajn = False
has_borisov = False
has_zhou_zhihua = False
has_zou_minghui = False
for p in doc2.paragraphs:
    text = p.text
    if 'Grinsztajn' in text:
        has_grinsztajn = True
    if 'Borisov' in text:
        has_borisov = True
    if '周志华' in text:
        has_zhou_zhihua = True
    if '邹明辉' in text:
        has_zou_minghui = True

print(f'Grinsztajn残留: {has_grinsztajn}')
print(f'Borisov残留: {has_borisov}')
print(f'周志华残留: {has_zhou_zhihua}')
print(f'邹明辉残留: {has_zou_minghui}')
